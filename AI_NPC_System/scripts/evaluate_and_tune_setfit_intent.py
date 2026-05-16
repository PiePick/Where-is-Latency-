"""Evaluate and tune the SWDA SetFit intent classifier for FastTrack.

The script uses the preprocessed SWDA coarse-label JSONL file, builds
leakage-aware validation/test splits, evaluates the current prototype model,
trains stronger balanced candidates, writes an Excel report, and saves the
best tuned model for later runtime integration.
"""

from __future__ import annotations

import argparse
import json
import random
import shutil
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from sklearn.metrics import accuracy_score, classification_report, confusion_matrix, f1_score


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DATA_PATH = ROOT / "prepared_fasttrack_data" / "swda_intent_coarse.jsonl"
DEFAULT_BASELINE_MODEL = ROOT / "models" / "setfit_swda_intent_minilm"
DEFAULT_OPTIMIZED_MODEL = ROOT / "models" / "setfit_swda_intent_minilm_optimized"
DEFAULT_RUN_DIR = ROOT / "models" / "setfit_swda_tuning_runs"
DEFAULT_REPORT_DIR = ROOT / "reports"
DEFAULT_BASE_MODEL = "sentence-transformers/all-MiniLM-L6-v2"
LABELS = ["QUESTION", "INFORM", "ACKNOWLEDGE", "DIRECTIVE", "EXPRESSIVE", "REJECT"]


@dataclass(frozen=True)
class CandidateConfig:
    """SetFit candidate training configuration."""

    name: str
    train_per_label: int
    num_iterations: int
    num_epochs: int
    batch_size: int


def parse_args() -> argparse.Namespace:
    """Parse CLI options."""
    parser = argparse.ArgumentParser(description="Evaluate and tune SetFit intent model.")
    parser.add_argument("--data-path", type=Path, default=DEFAULT_DATA_PATH)
    parser.add_argument("--baseline-model", type=Path, default=DEFAULT_BASELINE_MODEL)
    parser.add_argument("--optimized-model", type=Path, default=DEFAULT_OPTIMIZED_MODEL)
    parser.add_argument("--run-dir", type=Path, default=DEFAULT_RUN_DIR)
    parser.add_argument("--report-dir", type=Path, default=DEFAULT_REPORT_DIR)
    parser.add_argument("--base-model", default=DEFAULT_BASE_MODEL)
    parser.add_argument("--seed", type=int, default=20260517)
    parser.add_argument("--holdout-per-label", type=int, default=72)
    parser.add_argument("--natural-test-frac", type=float, default=0.15)
    parser.add_argument("--natural-test-max-per-label", type=int, default=2000)
    parser.add_argument("--device", default=None)
    parser.add_argument("--keep-runs", action="store_true")
    return parser.parse_args()


def load_rows(path: Path) -> pd.DataFrame:
    """Load preprocessed SWDA rows and normalize the required columns."""
    df = pd.read_json(path, lines=True)
    df = df[["text", "coarse_label", "source", "raw_dialog_act"]].copy()
    df["text"] = df["text"].astype(str).str.strip()
    df["coarse_label"] = df["coarse_label"].astype(str)
    df = df[df["text"].ne("") & df["coarse_label"].isin(LABELS)]
    return df.drop_duplicates(subset=["text", "coarse_label"]).reset_index(drop=True)


def load_leakage_keys(model_dir: Path) -> set[tuple[str, str]]:
    """Read prior training rows so evaluation splits do not reuse them."""
    train_path = model_dir / "swda_fewshot_train.jsonl"
    if not train_path.exists():
        return set()
    keys: set[tuple[str, str]] = set()
    with train_path.open("r", encoding="utf-8") as stream:
        for line in stream:
            if not line.strip():
                continue
            row = json.loads(line)
            keys.add((str(row.get("text", "")).strip(), str(row.get("coarse_label", "")).strip()))
    return keys


def build_splits(
    df: pd.DataFrame,
    *,
    leakage_keys: set[tuple[str, str]],
    seed: int,
    holdout_per_label: int,
    natural_test_frac: float,
    natural_test_max_per_label: int,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Build train pool, balanced validation/test, and natural test splits."""
    rng = random.Random(seed)
    rows = df.to_dict("records")
    rows = [
        row
        for row in rows
        if (str(row["text"]).strip(), str(row["coarse_label"]).strip()) not in leakage_keys
    ]

    train_parts: list[pd.DataFrame] = []
    val_parts: list[pd.DataFrame] = []
    test_parts: list[pd.DataFrame] = []
    natural_parts: list[pd.DataFrame] = []

    for label in LABELS:
        bucket = [row for row in rows if row["coarse_label"] == label]
        rng.shuffle(bucket)
        n_balanced = min(holdout_per_label, max(1, (len(bucket) - 1) // 4))
        val_rows = bucket[:n_balanced]
        test_rows = bucket[n_balanced : n_balanced * 2]
        remaining = bucket[n_balanced * 2 :]

        n_natural = min(
            natural_test_max_per_label,
            max(1, int(round(len(remaining) * natural_test_frac))),
        )
        natural_rows = remaining[:n_natural]
        train_rows = remaining[n_natural:]

        val_parts.append(pd.DataFrame(val_rows))
        test_parts.append(pd.DataFrame(test_rows))
        natural_parts.append(pd.DataFrame(natural_rows))
        train_parts.append(pd.DataFrame(train_rows))

    train_pool = pd.concat(train_parts, ignore_index=True).sample(frac=1, random_state=seed)
    val_balanced = pd.concat(val_parts, ignore_index=True).sample(frac=1, random_state=seed)
    test_balanced = pd.concat(test_parts, ignore_index=True).sample(frac=1, random_state=seed)
    test_natural = pd.concat(natural_parts, ignore_index=True).sample(frac=1, random_state=seed)
    return train_pool.reset_index(drop=True), val_balanced.reset_index(drop=True), test_balanced.reset_index(drop=True), test_natural.reset_index(drop=True)


def sample_train_data(train_pool: pd.DataFrame, train_per_label: int, seed: int) -> pd.DataFrame:
    """Build a balanced training sample for one candidate run."""
    parts = []
    for index, label in enumerate(LABELS):
        bucket = train_pool[train_pool["coarse_label"] == label]
        if bucket.empty:
            raise ValueError(f"No train rows available for label: {label}")
        n = min(train_per_label, len(bucket))
        parts.append(bucket.sample(n=n, random_state=seed + index))
    return pd.concat(parts, ignore_index=True).sample(frac=1, random_state=seed).reset_index(drop=True)


def to_setfit_dataset(df: pd.DataFrame) -> Any:
    """Convert a label-name DataFrame into a SetFit Dataset."""
    from datasets import Dataset

    label_to_id = {label: index for index, label in enumerate(LABELS)}
    records = [
        {"text": row["text"], "label": label_to_id[row["coarse_label"]]}
        for _, row in df.iterrows()
    ]
    return Dataset.from_list(records)


def predict_labels(model: Any, texts: list[str], *, batch_size: int = 256) -> list[str]:
    """Predict labels in batches and normalize integer IDs to label names."""
    predictions: list[str] = []
    for start in range(0, len(texts), batch_size):
        batch = texts[start : start + batch_size]
        raw = model.predict(batch)
        for item in raw:
            try:
                predictions.append(LABELS[int(item)])
            except Exception:
                predictions.append(str(item))
    return predictions


def evaluate_model(model: Any, df: pd.DataFrame, split_name: str) -> tuple[dict[str, Any], pd.DataFrame, pd.DataFrame]:
    """Evaluate a SetFit model and return summary, per-class rows, and confusions."""
    y_true = df["coarse_label"].tolist()
    y_pred = predict_labels(model, df["text"].tolist())
    report = classification_report(y_true, y_pred, labels=LABELS, output_dict=True, zero_division=0)
    matrix = confusion_matrix(y_true, y_pred, labels=LABELS)

    summary = {
        "split": split_name,
        "samples": len(df),
        "accuracy": float(accuracy_score(y_true, y_pred)),
        "macro_f1": float(f1_score(y_true, y_pred, labels=LABELS, average="macro", zero_division=0)),
        "weighted_f1": float(f1_score(y_true, y_pred, labels=LABELS, average="weighted", zero_division=0)),
    }

    per_class_rows = []
    for label in LABELS:
        stats = report[label]
        per_class_rows.append(
            {
                "split": split_name,
                "label": label,
                "precision": float(stats["precision"]),
                "recall": float(stats["recall"]),
                "f1": float(stats["f1-score"]),
                "support": int(stats["support"]),
            }
        )

    confusion_rows = []
    for true_index, true_label in enumerate(LABELS):
        for pred_index, pred_label in enumerate(LABELS):
            count = int(matrix[true_index, pred_index])
            if true_label != pred_label and count:
                confusion_rows.append(
                    {
                        "split": split_name,
                        "true_label": true_label,
                        "predicted_label": pred_label,
                        "count": count,
                    }
                )
    return summary, pd.DataFrame(per_class_rows), pd.DataFrame(confusion_rows)


def train_candidate(
    train_df: pd.DataFrame,
    config: CandidateConfig,
    *,
    output_dir: Path,
    base_model: str,
) -> tuple[Any, float]:
    """Train one SetFit candidate and save it."""
    from setfit import SetFitModel, Trainer, TrainingArguments

    started = time.perf_counter()
    if output_dir.exists():
        shutil.rmtree(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    dataset = to_setfit_dataset(train_df)
    model = SetFitModel.from_pretrained(base_model, labels=LABELS)
    args = TrainingArguments(
        output_dir=str(output_dir / "trainer"),
        batch_size=config.batch_size,
        num_epochs=config.num_epochs,
        num_iterations=config.num_iterations,
    )
    trainer = Trainer(
        model=model,
        args=args,
        train_dataset=dataset,
        column_mapping={"text": "text", "label": "label"},
    )
    trainer.train()
    model.save_pretrained(str(output_dir))

    metadata = {
        "base_model": base_model,
        "candidate": config.__dict__,
        "labels": LABELS,
        "train_size": int(len(train_df)),
        "label_counts": train_df["coarse_label"].value_counts().to_dict(),
    }
    (output_dir / "fasttrack_setfit_metadata.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    train_df.to_json(output_dir / "swda_fewshot_train.jsonl", orient="records", lines=True, force_ascii=False)
    shutil.rmtree(output_dir / "trainer", ignore_errors=True)
    return model, time.perf_counter() - started


def write_excel_report(
    path: Path,
    *,
    summary_df: pd.DataFrame,
    per_class_df: pd.DataFrame,
    confusion_df: pd.DataFrame,
    split_counts_df: pd.DataFrame,
    best_matrix_df: pd.DataFrame,
    notes: list[str],
) -> None:
    """Write a formatted Excel evaluation workbook."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary_df.to_excel(writer, index=False, sheet_name="Summary")
        split_counts_df.to_excel(writer, index=False, sheet_name="Dataset Splits")
        per_class_df.to_excel(writer, index=False, sheet_name="Per Class")
        confusion_df.to_excel(writer, index=False, sheet_name="Confusions")
        best_matrix_df.to_excel(writer, index=True, sheet_name="Best Confusion Matrix")
        pd.DataFrame({"note": notes}).to_excel(writer, index=False, sheet_name="Notes")

    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for column_cells in sheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in column_cells)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max_len + 2, 42)
    workbook.save(path)


def main() -> int:
    """Run evaluation, tuning, and report generation."""
    args = parse_args()
    df = load_rows(args.data_path)
    leakage_keys = load_leakage_keys(args.baseline_model)
    train_pool, val_balanced, test_balanced, test_natural = build_splits(
        df,
        leakage_keys=leakage_keys,
        seed=args.seed,
        holdout_per_label=args.holdout_per_label,
        natural_test_frac=args.natural_test_frac,
        natural_test_max_per_label=args.natural_test_max_per_label,
    )

    from setfit import SetFitModel

    candidates = [
        CandidateConfig("optimized_240_iter6", 240, 6, 1, 16),
        CandidateConfig("optimized_320_iter8", 320, 8, 1, 16),
        CandidateConfig("optimized_338_iter10", 338, 10, 1, 16),
    ]

    summary_rows: list[dict[str, Any]] = []
    per_class_frames: list[pd.DataFrame] = []
    confusion_frames: list[pd.DataFrame] = []
    best_name = "baseline_1080"
    best_score = -1.0
    best_model_dir = args.baseline_model
    best_matrix_source: tuple[Any, pd.DataFrame] | None = None

    baseline = SetFitModel.from_pretrained(str(args.baseline_model))
    for split_name, split_df in (("balanced_val", val_balanced), ("balanced_test", test_balanced), ("natural_test", test_natural)):
        summary, per_class, confusions = evaluate_model(baseline, split_df, split_name)
        summary.update({"model": "baseline_1080", "train_seconds": 0.0, "train_per_label": 180, "num_iterations": 5})
        summary_rows.append(summary)
        per_class["model"] = "baseline_1080"
        confusions["model"] = "baseline_1080"
        per_class_frames.append(per_class)
        confusion_frames.append(confusions)
        if split_name == "balanced_val":
            best_score = summary["macro_f1"]

    for config in candidates:
        train_df = sample_train_data(train_pool, config.train_per_label, args.seed)
        candidate_dir = args.run_dir / config.name
        model, train_seconds = train_candidate(
            train_df,
            config,
            output_dir=candidate_dir,
            base_model=args.base_model,
        )
        val_summary, val_per_class, val_confusions = evaluate_model(model, val_balanced, "balanced_val")
        test_summary, test_per_class, test_confusions = evaluate_model(model, test_balanced, "balanced_test")

        for summary, per_class, confusions in (
            (val_summary, val_per_class, val_confusions),
            (test_summary, test_per_class, test_confusions),
        ):
            summary.update(
                {
                    "model": config.name,
                    "train_seconds": round(train_seconds, 3),
                    "train_per_label": config.train_per_label,
                    "num_iterations": config.num_iterations,
                }
            )
            summary_rows.append(summary)
            per_class["model"] = config.name
            confusions["model"] = config.name
            per_class_frames.append(per_class)
            confusion_frames.append(confusions)

        if val_summary["macro_f1"] > best_score:
            best_score = val_summary["macro_f1"]
            best_name = config.name
            best_model_dir = candidate_dir
            best_matrix_source = (model, test_balanced)

    if best_model_dir != args.baseline_model:
        if args.optimized_model.exists():
            shutil.rmtree(args.optimized_model)
        shutil.copytree(best_model_dir, args.optimized_model)
        best_model_dir = args.optimized_model
    else:
        if args.optimized_model.exists():
            shutil.rmtree(args.optimized_model)
        shutil.copytree(args.baseline_model, args.optimized_model)
        baseline = SetFitModel.from_pretrained(str(args.optimized_model))
        best_matrix_source = (baseline, test_balanced)

    best_model = SetFitModel.from_pretrained(str(args.optimized_model))
    natural_summary, natural_per_class, natural_confusions = evaluate_model(best_model, test_natural, "natural_test")
    natural_summary.update(
        {
            "model": f"{best_name}_selected",
            "train_seconds": np.nan,
            "train_per_label": np.nan,
            "num_iterations": np.nan,
        }
    )
    summary_rows.append(natural_summary)
    natural_per_class["model"] = f"{best_name}_selected"
    natural_confusions["model"] = f"{best_name}_selected"
    per_class_frames.append(natural_per_class)
    confusion_frames.append(natural_confusions)

    y_true = test_balanced["coarse_label"].tolist()
    y_pred = predict_labels(best_model, test_balanced["text"].tolist())
    best_matrix = confusion_matrix(y_true, y_pred, labels=LABELS)
    best_matrix_df = pd.DataFrame(best_matrix, index=[f"true_{x}" for x in LABELS], columns=[f"pred_{x}" for x in LABELS])

    summary_df = pd.DataFrame(summary_rows)
    split_counts = []
    for split_name, split_df in (
        ("train_pool", train_pool),
        ("balanced_val", val_balanced),
        ("balanced_test", test_balanced),
        ("natural_test", test_natural),
    ):
        counts = split_df["coarse_label"].value_counts().to_dict()
        row = {"split": split_name, "samples": len(split_df)}
        row.update({label: counts.get(label, 0) for label in LABELS})
        split_counts.append(row)
    split_counts_df = pd.DataFrame(split_counts)
    per_class_df = pd.concat(per_class_frames, ignore_index=True)
    confusion_df = pd.concat(confusion_frames, ignore_index=True)
    if not confusion_df.empty:
        confusion_df = confusion_df.sort_values(["model", "split", "count"], ascending=[True, True, False])

    args.report_dir.mkdir(parents=True, exist_ok=True)
    json_path = args.report_dir / "setfit_intent_evaluation_summary.json"
    xlsx_path = args.report_dir / "setfit_intent_evaluation.xlsx"
    notes = [
        "Balanced validation/test splits exclude rows used by the baseline 1,080-sample prototype when available.",
        "Candidate selection uses balanced validation macro-F1 to avoid hiding minority-label failures behind INFORM/ACKNOWLEDGE dominance.",
        "Natural test is capped per label for runtime while preserving the corpus label skew.",
        f"Selected model: {best_name}; saved to {args.optimized_model}",
    ]
    write_excel_report(
        xlsx_path,
        summary_df=summary_df,
        per_class_df=per_class_df,
        confusion_df=confusion_df,
        split_counts_df=split_counts_df,
        best_matrix_df=best_matrix_df,
        notes=notes,
    )
    json_path.write_text(
        json.dumps(
            {
                "selected_model": best_name,
                "optimized_model_dir": str(args.optimized_model),
                "summary": summary_df.to_dict(orient="records"),
                "split_counts": split_counts_df.to_dict(orient="records"),
                "top_confusions": confusion_df.head(30).to_dict(orient="records"),
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    if not args.keep_runs:
        shutil.rmtree(args.run_dir, ignore_errors=True)

    print(json.dumps({
        "selected_model": best_name,
        "best_validation_macro_f1": round(float(best_score), 6),
        "optimized_model_dir": str(args.optimized_model),
        "excel_report": str(xlsx_path),
        "json_report": str(json_path),
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
